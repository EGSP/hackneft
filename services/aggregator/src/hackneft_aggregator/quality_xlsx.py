"""Разбор выгрузок качества (ЛИМС и ПАК) из xlsx.

Показатель занимает пару столбцов «время — значение». Ряды независимы и разной длины, поэтому
строка книги моментом времени не является. Для единообразия с телеметрией агрегатор читает уже
переложенные в CSV длинного формата (`date,tag,value`) файлы; этот модуль — только разбор xlsx
при конвертации.
"""

from __future__ import annotations

import re
from collections.abc import Callable, Iterator, Sequence
from dataclasses import dataclass
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

HeaderPair = tuple[tuple[str, str], ...]
TagNamer = Callable[[Sequence[HeaderPair]], Sequence[str | None]]

SERVICE_HEADINGS = frozenset(
    {"время", "дата", "значение", "результат", "time", "date", "value", "result"}
)

LIMS_RESULT_DELAY = timedelta(hours=4)
"""Задержка появления лабораторного результата после отбора пробы (по спецификации данных)."""

_INSTALLATIONS = (
    ("гидроочистк", "ht"),
    ("авт", "avt"),
)
_POINT_NUMBER = re.compile(
    r"точка\s+отбора\s*№?\s*['\"]?(\d+(?:\.\d+)*)['\"]?",
    re.IGNORECASE,
)
_SLUG_SEPARATORS = re.compile(r"[^0-9a-zа-яё]+", re.IGNORECASE)


@dataclass(frozen=True, slots=True)
class QualityReading:
    timestamp: datetime
    tag: str
    value: float


@dataclass(frozen=True, slots=True)
class _Pair:
    tag: str
    time_index: int
    value_index: int


def convert_pak_xlsx(path: Path) -> list[QualityReading]:
    """Читает выгрузку ПАК. Имена тегов — из шапки как есть (`24-2000:Mg.Sulfur`).

    Между парами столбцов бывает пустой разделитель, поэтому пары ищутся по ячейкам с
    именами тегов, а не жёстким шагом 2. Вторая строка шапки — единицы измерения.
    """
    return _collect(
        path,
        header_rows=2,
        tag_row=0,
        name_tags=_pak_names,
        delay=timedelta(0),
    )


def convert_lims_xlsx(path: Path) -> list[QualityReading]:
    """Читает выгрузку ЛИМС.

    Имена тегов вида `ht.2.Mg.Sulfur` (установка, точка, показатель). В CSV время — момент
    доступности результата (отбор + 4 часа), а не момент отбора пробы.
    """
    return _collect(
        path,
        header_rows=4,
        tag_row=1,
        name_tags=_lims_names,
        delay=LIMS_RESULT_DELAY,
    )


def write_long_csv(path: Path, readings: Sequence[QualityReading]) -> None:
    """Пишет длинный CSV: `date,tag,value`, строки упорядочены по времени и тегу."""
    ordered = sorted(readings, key=lambda item: (item.timestamp, item.tag))
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as text:
        text.write("date,tag,value\n")
        for reading in ordered:
            stamp = reading.timestamp.isoformat(sep=" ", timespec="seconds")
            text.write(f"{stamp},{reading.tag},{_format_value(reading.value)}\n")


def _collect(
    path: Path,
    *,
    header_rows: int,
    tag_row: int,
    name_tags: TagNamer,
    delay: timedelta,
) -> list[QualityReading]:
    workbook = load_workbook(filename=path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        if sheet is None:
            raise ValueError(f"книга пуста: {path}")
        rows = sheet.iter_rows(values_only=True)
        header = _read_header(rows, header_rows)
        pairs = _pairs(header, name_tags, tag_row=tag_row)
        if not pairs:
            raise ValueError(f"в шапке не нашлось ни одной пары: {path}")

        collected: list[QualityReading] = []
        for row in rows:
            for pair in pairs:
                reading = _reading(pair, row, delay)
                if reading is not None:
                    collected.append(reading)
        return collected
    finally:
        workbook.close()


def _read_header(rows: Iterator[tuple[Any, ...]], header_rows: int) -> list[list[str]]:
    header: list[list[str]] = []
    for _ in range(header_rows):
        row = next(rows, None)
        if row is None:
            raise ValueError("в выгрузке нет шапки")
        header.append([_text(cell) for cell in row])
    return header


def _pairs(header: list[list[str]], name_tags: TagNamer, *, tag_row: int) -> list[_Pair]:
    """Пары столбцов начинаются там, где в строке имён тега есть непустое имя.

    Так находятся и соседние пары ЛИМС (шаг 2), и пары ПАК с пустым столбцом между ними.
    """
    if tag_row >= len(header):
        raise ValueError(f"строка имён тегов {tag_row} отсутствует в шапке")
    name_cells = header[tag_row]
    columns = [
        index
        for index, cell in enumerate(name_cells)
        if cell.strip() != "" and not _is_service_heading(cell)
    ]
    headings: list[HeaderPair] = [
        tuple((_at(row, left), _at(row, left + 1)) for row in header) for left in columns
    ]
    names = name_tags(headings)
    pairs: list[_Pair] = []
    seen: set[str] = set()
    for left, tag in zip(columns, names, strict=True):
        if tag is None or tag in seen:
            continue
        seen.add(tag)
        pairs.append(_Pair(tag=tag, time_index=left, value_index=left + 1))
    return pairs


def _pak_names(headings: Sequence[HeaderPair]) -> list[str | None]:
    return [_first_meaningful(pair[0]) or None for pair in headings]


def _lims_names(headings: Sequence[HeaderPair]) -> list[str | None]:
    """Точка отбора объединена над своими показателями — пустая ячейка наследует предыдущую."""
    names: list[str | None] = []
    point = ""
    for pair in headings:
        point = _first_meaningful(pair[0]) or point
        indicator = _first_meaningful(pair[1]) if len(pair) > 1 else ""
        names.append(_lims_tag(point, indicator))
    return names


def _lims_tag(point: str, indicator: str) -> str | None:
    if indicator.strip() == "":
        return None
    place = _place(point) or _slug(point)
    if place == "":
        return indicator.strip()
    return f"{place}.{indicator.strip()}"


def _place(point: str) -> str:
    lowered = point.lower()
    installation = next(
        (code for marker, code in _INSTALLATIONS if marker in lowered),
        "",
    )
    number = _POINT_NUMBER.search(lowered)
    if installation == "" or number is None:
        return ""
    return f"{installation}.{number.group(1)}"


def _slug(text: str) -> str:
    return _SLUG_SEPARATORS.sub("-", text.strip().lower()).strip("-")


def _is_service_heading(text: str) -> bool:
    return text.strip().lower().rstrip(":") in SERVICE_HEADINGS


def _first_meaningful(cells: Sequence[str]) -> str:
    for cell in cells:
        text = cell.strip()
        if text != "" and not _is_service_heading(text):
            return text
    return ""


def _reading(pair: _Pair, row: Sequence[Any], delay: timedelta) -> QualityReading | None:
    raw_time = _cell(row, pair.time_index)
    raw_value = _cell(row, pair.value_index)
    if raw_time is None and raw_value is None:
        return None
    timestamp = _parse_timestamp(raw_time)
    if timestamp is None:
        return None
    value = _parse_number(raw_value)
    if value is None:
        return None
    return QualityReading(timestamp=timestamp + delay, tag=pair.tag, value=value)


def _cell(row: Sequence[Any], index: int) -> Any:
    if index >= len(row):
        return None
    value = row[index]
    if isinstance(value, str) and value.strip() == "":
        return None
    return value


def _at(row: Sequence[str], index: int) -> str:
    return row[index] if index < len(row) else ""


def _text(cell: Any) -> str:
    if cell is None:
        return ""
    if isinstance(cell, datetime):
        return cell.isoformat()
    return str(cell)


def _parse_timestamp(value: Any) -> datetime | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    text = str(value).strip().replace(" ", "T", 1)
    try:
        return datetime.fromisoformat(text)
    except ValueError:
        return None


def _parse_number(value: Any) -> float | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        number = float(value)
        return number if number == number and number not in (float("inf"), float("-inf")) else None
    text = str(value).strip().replace(",", ".")
    if text == "":
        return None
    try:
        number = float(text)
    except ValueError:
        return None
    return number if number == number and number not in (float("inf"), float("-inf")) else None


def _format_value(value: float) -> str:
    text = f"{value:.12g}"
    return text
